import os
from pathlib import Path
import fitz # PyMuPDF
import docx
from PIL import Image, ImageDraw, ImageFont

FIXTURES_DIR = Path(__file__).parent / "fixtures"

def generate_all_fixtures(force: bool = False):
    os.makedirs(FIXTURES_DIR, exist_ok=True)
    if not force and (FIXTURES_DIR / "sample_multipage_skew.pdf").exists() and (FIXTURES_DIR / "sample_rrhh_payroll.pdf").exists():
        return
    generate_hr_pdf()
    generate_legal_docx()
    generate_support_pdf()
    generate_scanned_raster_pdf()
    generate_rotated_scanned_pdf()
    generate_skewed_fixtures()
    print("All synthetic fixtures (including deskew suite) generated successfully.")

def _draw_sample_scanned_image(extra_skew: float = 0.0, add_edge_entities: bool = False) -> Image.Image:
    width, height = 1240, 1754 # ~150 DPI A4
    img = Image.new("RGB", (width, height), color="white")
    draw = ImageDraw.Draw(img)

    try:
        font_title = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 36)
        font_body = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 26)
    except Exception:
        font_title = ImageFont.load_default()
        font_body = ImageFont.load_default()

    draw.text((100, 100), "INFORME MEDICO DE APTITUD LABORAL — ESCANEADO", fill=(20, 20, 20), font=font_title)
    draw.text((100, 160), "CENTRO MEDICO LABORAL ANCLORA — DEPARTAMENTO DE RRHH", fill=(80, 80, 80), font=font_body)
    draw.line([(100, 210), (1140, 210)], fill=(120, 120, 120), width=3)

    items = [
        ("Trabajador:", "Maria Dolores Santos Ruiz"),
        ("DNI / NIE:", "54321987M"),
        ("Email Contacto:", "maria.santos@example.test"),
        ("Telefono Movil:", "+34 633 445 566"),
        ("Cuenta Bancaria:", "ES12 3456 7890 1234 5678 9012"),
        ("Aptitud Medica:", "APTO PARA EL PUESTO DE TRABAJO"),
        ("Notas Confidenciales:", "Sin patologias previas de riesgo laboral.")
    ]

    y = 260
    for label, val in items:
        draw.text((100, y), label, fill=(50, 50, 50), font=font_body)
        draw.text((450, y), val, fill=(0, 0, 0), font=font_body)
        y += 65

    if add_edge_entities:
        # Entity near left/top border and bottom border
        draw.text((25, 40), "REF-BORDER-DNI: 77889900X", fill=(0, 0, 0), font=font_body)
        draw.text((25, 1680), "CONFIDENCIAL-IBAN: ES88 9900 1122 3344 5566 7788", fill=(0, 0, 0), font=font_body)

    if abs(extra_skew) > 0.01:
        # Rotate bitmap image to simulate physical feeder skew
        # expand=False keeps page size, fillcolor white
        img = img.rotate(extra_skew, resample=Image.BICUBIC, expand=False, fillcolor="white")

    return img

def generate_skewed_fixtures():
    # 1. Tilted +2 degrees
    img_plus_2 = _draw_sample_scanned_image(extra_skew=2.0)
    _save_image_as_pdf(img_plus_2, FIXTURES_DIR / "sample_skew_plus_2deg.pdf")

    # 2. Tilted -5 degrees
    img_minus_5 = _draw_sample_scanned_image(extra_skew=-5.0)
    _save_image_as_pdf(img_minus_5, FIXTURES_DIR / "sample_skew_minus_5deg.pdf")

    # 3. Tilted +12 degrees
    img_plus_12 = _draw_sample_scanned_image(extra_skew=12.0)
    _save_image_as_pdf(img_plus_12, FIXTURES_DIR / "sample_skew_plus_12deg.pdf")

    # 4. Rotated 90 deg + 3 deg skew
    img_rot_skew = _draw_sample_scanned_image(extra_skew=3.0)
    _save_image_as_pdf(img_rot_skew, FIXTURES_DIR / "sample_rot90_skew_3deg.pdf", rotation=90)

    # 5. Multipage document with different angles
    doc_multi = fitz.open()
    for angle in [0.0, 3.5, -4.0]:
        img_page = _draw_sample_scanned_image(extra_skew=angle)
        tmp_p = FIXTURES_DIR / f"tmp_multi_{os.getpid()}_{angle}.png"
        img_page.save(str(tmp_p), format="PNG")
        page = doc_multi.new_page(width=595, height=842)
        page.insert_image(page.rect, filename=str(tmp_p))
        if tmp_p.exists():
            tmp_p.unlink()
    doc_multi.save(str(FIXTURES_DIR / "sample_multipage_skew.pdf"))
    doc_multi.close()

    # 6. Sensitive entities near edges
    img_edge = _draw_sample_scanned_image(extra_skew=2.5, add_edge_entities=True)
    _save_image_as_pdf(img_edge, FIXTURES_DIR / "sample_edge_entities_skew.pdf")

    # 7. Low confidence / ambiguous page (no text lines)
    img_blank = Image.new("RGB", (1240, 1754), color="white")
    draw = ImageDraw.Draw(img_blank)
    draw.rectangle([500, 500, 700, 700], fill="blue") # Single square block, no lines
    _save_image_as_pdf(img_blank, FIXTURES_DIR / "sample_low_confidence_unskewable.pdf")

def _save_image_as_pdf(img: Image.Image, out_path: Path, rotation: int = 0):
    import os as _os, uuid as _uuid
    tmp_path = out_path.with_suffix(f".tmp.{_os.getpid()}.{_uuid.uuid4().hex[:8]}.png")
    img.save(str(tmp_path), format="PNG")
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, filename=str(tmp_path))
    if rotation:
        page.set_rotation(rotation)
    doc.save(str(out_path))
    doc.close()
    if tmp_path.exists():
        tmp_path.unlink()

def generate_hr_pdf():
    pdf_path = FIXTURES_DIR / "sample_rrhh_payroll.pdf"
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 60), "RECIBO DE NOMINA CONFIDENCIAL — DICIEMBRE 2025", fontsize=14, fontname="helv", color=(0.1, 0.2, 0.5))
    page.insert_text((50, 85), "EMPRESA: Soluciones Tecnologicas Ibericas S.L.", fontsize=10, fontname="helv")
    
    y = 120
    lines = [
        ("Nombre del Empleado:", "Laura Martinez Gomez"),
        ("DNI / Identificacion Fiscal:", "12345678Z"),
        ("Numero Afiliacion Seguridad Social:", "NSS 28 12345678 40"),
        ("Direccion Postal:", "Calle Mayor 42, 28013 Madrid"),
        ("Correo Electronico Corporativo:", "laura.martinez@example.test"),
        ("Telefono de Contacto:", "+34 600 123 456"),
        ("Cuenta Bancaria (IBAN):", "ES00 0000 0000 0000 0000 0000"),
        ("Categoria Profesional:", "Ingeniera de Software Senior"),
        ("Salario Base Mensual:", "3.850,00 EUR"),
        ("Complemento de Destino:", "450,00 EUR"),
        ("Total Liquido a Percibir:", "4.300,00 EUR")
    ]
    for label, val in lines:
        page.insert_text((50, y), label, fontsize=10, fontname="helv", color=(0.2, 0.2, 0.2))
        page.insert_text((240, y), val, fontsize=10, fontname="helv", color=(0, 0, 0))
        y += 24
        
    page.insert_text((50, 480), "Documento de caracter estrictamente confidencial bajo normativa RGPD.", fontsize=9, fontname="helv", color=(0.4, 0.4, 0.4))
    doc.set_metadata({
        "author": "Laura Martinez Gomez",
        "creator": "Payroll System v2",
        "subject": "Nomina 12345678Z"
    })
    doc.save(str(pdf_path))
    doc.close()
    return str(pdf_path)

def generate_legal_docx():
    docx_path = FIXTURES_DIR / "sample_legal_contract.docx"
    doc = docx.Document()
    section = doc.sections[0]
    header = section.header
    header.paragraphs[0].text = "ACUERDO DE CONFIDENCIALIDAD Y PRESTACION DE SERVICIOS (REF: EXP-9982/2025)"
    doc.add_heading("CONTRATO DE ARRENDAMIENTO DE SERVICIOS", level=1)
    
    doc.add_paragraph("En Madrid, a 15 de enero de 2026. COMPARECEN:")
    doc.add_paragraph(
        "De una parte, Don Carlos Fernandez Soto, mayor de edad, con DNI 87654321B, con domicilio profesional en "
        "Paseo de la Castellana 100, 28046 Madrid, y correo carlos.fernandez@example.test, en nombre y representacion de "
        "IBERICA CONSULTING S.A., con CIF A12345678."
    )
    doc.add_paragraph(
        "De otra parte, Dona Elena Ramos Blanco, con DNI 44556677C, en nombre de TECH SOLUTIONS SL, con CIF B98765432, "
        "con telefono +34 654 987 321 y correo elena.ramos@example.test."
    )
    doc.add_heading("CLAUSULAS", level=2)
    doc.add_paragraph(
        "PRIMERA. - Objeto del contrato: Prestacion de asesoria legal en el Procedimiento Judicial nº 452/2025 "
        "ante el Juzgado de Primera Instancia nº 4 de Madrid."
    )
    doc.add_paragraph(
        "SEGUNDA. - Honorarios: Se fija un importe total de 12.500,00 EUR que sera abonado mediante transferencia "
        "a la cuenta bancaria ES99 1234 5678 9012 3456 7890."
    )
    
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Representante"
    hdr_cells[1].text = "DNI"
    hdr_cells[2].text = "Email"
    
    row_cells = table.add_row().cells
    row_cells[0].text = "Carlos Fernandez Soto"
    row_cells[1].text = "87654321B"
    row_cells[2].text = "carlos.fernandez@example.test"
    
    row_cells2 = table.add_row().cells
    row_cells2[0].text = "Elena Ramos Blanco"
    row_cells2[1].text = "44556677C"
    row_cells2[2].text = "elena.ramos@example.test"
    
    doc.core_properties.author = "Carlos Fernandez Soto"
    doc.core_properties.comments = "Borrador de contrato confidencial"
    doc.save(str(docx_path))
    return str(docx_path)

def generate_support_pdf():
    pdf_path = FIXTURES_DIR / "sample_soporte_incident.pdf"
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((50, 60), "INFORME DE INCIDENCIA TECNICA Y DIAGNOSTICO DE RED", fontsize=14, fontname="helv", color=(0.8, 0.1, 0.1))
    page.insert_text((50, 85), "DEVOPS & CYBERSECURITY OPERATIONS TICKET #8841", fontsize=10, fontname="helv")
    
    y = 120
    lines = [
        ("Ingeniero de Guardia:", "David Navarro Ruiz"),
        ("Email de Contacto:", "david.navarro@example.test"),
        ("Telefono Urgencias:", "+34 611 223 344"),
        ("Identificador Cliente:", "CID-9942"),
        ("Direccion IP Afectada (Host 1):", "192.168.1.105"),
        ("Direccion IP Bastion Publico:", "198.51.100.45"),
        ("API Token de Integracion:", "purgedoc_test_token_REDACTED"),
        ("Token de Despliegue CI/CD:", "ghp_1234567890abcdefghijklmnopqrstuvwxyz"),
        ("Acceso a Base de Datos:", "Host: 10.0.4.12:5432 con usuario admin"),
        ("Diagnostico del Incidente:", "Fuga potencial de credenciales mitigada temporalmente.")
    ]
    for label, val in lines:
        page.insert_text((50, y), label, fontsize=10, fontname="helv", color=(0.2, 0.2, 0.2))
        page.insert_text((230, y), val, fontsize=10, fontname="helv", color=(0, 0, 0))
        y += 24
        
    doc.set_metadata({
        "author": "David Navarro Ruiz",
        "creator": "Incident Reporter",
        "keywords": "198.51.100.45, token leak"
    })
    doc.save(str(pdf_path))
    doc.close()
    return str(pdf_path)

def generate_scanned_raster_pdf():
    pdf_path = FIXTURES_DIR / "sample_scanned_medical_hr.pdf"
    img = _draw_sample_scanned_image(extra_skew=0.0)
    _save_image_as_pdf(img, pdf_path)
    return str(pdf_path)

def generate_rotated_scanned_pdf():
    pdf_path = FIXTURES_DIR / "sample_rotated_scanned.pdf"
    img = _draw_sample_scanned_image(extra_skew=0.0)
    _save_image_as_pdf(img, pdf_path, rotation=90)
    return str(pdf_path)

if __name__ == "__main__":
    generate_all_fixtures()
