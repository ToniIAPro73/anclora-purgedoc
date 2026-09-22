import os
import shutil
import zipfile
import subprocess
import logging
from typing import Dict, Any, List, Tuple
from pathlib import Path
import fitz # PyMuPDF
import docx
from lxml import etree
from backend.models import DocumentMetadata, MatchItem, BoundingBox
from backend.services.ocr import local_ocr_engine

logger = logging.getLogger(__name__)

class DocumentProcessor:
    def __init__(self):
        pass

    def extract_pdf_content(
        self,
        pdf_path: str,
        apply_deskew: bool = True
    ) -> Tuple[List[Dict[str, Any]], int, bool, bool]:
        """
        Extracts text, page dimensions and bounding boxes using PyMuPDF.
        If the PDF has no selectable text layer (scanned/raster), seamlessly
        applies local Tesseract OCR with multi-angle deskew pre-pass.
        Returns:
            (pages_content, page_count, has_text_layer, is_scanned_ocr)
        """
        doc = fitz.open(pdf_path)
        page_count = len(doc)
        pages_content = []
        total_text_len = 0

        is_scanned = local_ocr_engine.is_raster_only_pdf(doc)

        for p_idx in range(page_count):
            page = doc[p_idx]
            page_w = page.rect.width
            page_h = page.rect.height
            deskew_info = {}

            if is_scanned:
                logger.info(f"Page {p_idx + 1} has no text layer. Running local Tesseract OCR with deskew...")
                p_text, rects, deskew_info = local_ocr_engine.ocr_page(page, p_idx + 1, apply_deskew=apply_deskew)
                total_text_len += len(p_text.strip())
            else:
                p_text = page.get_text("text")
                total_text_len += len(p_text.strip())
                words = page.get_text("words")
                rects = []
                for w in words:
                    rects.append({
                        "bbox": [w[0], w[1], w[2], w[3]],
                        "text": w[4]
                    })

            pages_content.append({
                "page_num": p_idx + 1,
                "text": p_text,
                "rects": rects,
                "width": page_w,
                "height": page_h,
                "is_ocr": is_scanned,
                "deskew": deskew_info
            })

        doc.close()
        has_text_layer = total_text_len > 0
        return pages_content, page_count, has_text_layer, is_scanned

    def convert_docx_to_preview_pdf(self, docx_path: str, output_dir: str) -> str:
        preview_pdf_path = os.path.join(output_dir, "preview.pdf")
        try:
            cmd = [
                "soffice",
                "--headless",
                "--convert-to", "pdf",
                "--outdir", output_dir,
                docx_path
            ]
            subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=True, timeout=30)
            
            base_name = Path(docx_path).stem
            converted_name = os.path.join(output_dir, f"{base_name}.pdf")
            if os.path.exists(converted_name):
                shutil.move(converted_name, preview_pdf_path)
                return preview_pdf_path
        except Exception as e:
            logger.warning(f"LibreOffice conversion failed, fallbacking to synthetic text PDF: {e}")
            
        self._create_fallback_docx_pdf(docx_path, preview_pdf_path)
        return preview_pdf_path

    def _create_fallback_docx_pdf(self, docx_path: str, out_pdf_path: str):
        doc = fitz.open()
        page = doc.new_page(width=595, height=842)
        doc_in = docx.Document(docx_path)
        y = 50
        for p in doc_in.paragraphs:
            if p.text:
                page.insert_text((50, y), p.text[:80], fontsize=11)
                y += 18
                if y > 800:
                    page = doc.new_page(width=595, height=842)
                    y = 50
        doc.save(out_pdf_path)
        doc.close()

    def extract_docx_content(self, docx_path: str) -> Tuple[List[Dict[str, Any]], int, bool]:
        doc = docx.Document(docx_path)
        full_text = []

        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        full_text.append(cell.text.strip())

        for section in doc.sections:
            for hp in section.header.paragraphs:
                if hp.text.strip():
                    full_text.append(hp.text.strip())
            for fp in section.footer.paragraphs:
                if fp.text.strip():
                    full_text.append(fp.text.strip())

        joined_text = "\n".join(full_text)
        has_text = len(joined_text.strip()) > 0

        pages_content = [{
            "page_num": 1,
            "text": joined_text,
            "rects": [],
            "width": 595.0,
            "height": 842.0
        }]
        return pages_content, 1, has_text

    def render_pdf_page_image(self, pdf_path: str, page_num: int) -> bytes:
        doc = fitz.open(pdf_path)
        p_idx = max(0, min(page_num - 1, len(doc) - 1))
        page = doc[p_idx]
        pix = page.get_pixmap(dpi=150)
        png_bytes = pix.tobytes("png")
        doc.close()
        return png_bytes

document_processor = DocumentProcessor()
